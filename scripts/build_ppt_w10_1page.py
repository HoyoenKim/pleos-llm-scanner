#!/usr/bin/env python
"""10주차 progress report — 1 page version.

같은 textbox layout (Rectangle 3, 16 paragraphs) 패턴이라 9주차 v3 빌드와
동일 update mechanism. 사용자가 미리 작성한 `pptx/자율주행연구프로젝트1_10주차_*.pptx`
를 base 로 사용하며 (textbox 구조 보존), paragraph 텍스트만 본 학기 측정값에
맞춰 갱신한다.

작성자/제출 파일명은 `scripts/_author.json` 에서 로드.
"""
import json
import shutil
from pathlib import Path

from pptx import Presentation


def _load_author() -> dict:
    here = Path(__file__).parent
    p = here / "_author.json"
    if not p.exists():
        p = here / "_author.example.json"
    with p.open(encoding="utf-8") as f:
        return json.load(f)


AUTHOR = _load_author()
ROOT = Path(__file__).parent.parent.parent
WEEK = "10"
PPT_BASENAME = f"자율주행연구프로젝트1_{WEEK}주차"
SRC = ROOT / "pptx" / f"{PPT_BASENAME}_{AUTHOR['student_id']}_{AUTHOR['author_name']}.pptx"
DST = ROOT / "pptx" / "v3" / f"{PPT_BASENAME}_{AUTHOR['student_id']}_{AUTHOR['author_name']}_v3.pptx"


# 16 paragraphs, indexed 0~15. 9주차 v3 와 동일 layout 가정.
NEW_TEXTS: dict[int, str | tuple[str, str]] = {
    # 0: Progress Report (period) — keep
    # 1: 프로젝트 제목 — keep
    2: "사례 연구 5건 + AAOS/TARA 매핑 + 본 학기 연구 RQ 4개 중 3개 첫 측정 + 보고서 v0.9-r1:",
    3: ("9주차까지 확정한 측정값(combined n=19, Stage 1 F1 0.882 / Stage 3 합의 ≥2/3 F1 0.966)을 보고서 v0.9-r1로 "
        "통합. 본 학기 연구 핵심 RQ 4개 중 3개에 corpus-level 첫 정량 인덱스 추가. 보고서 가독성 정제 + 외부 공개판 분리."),
    # 4: "1) " + "사례 연구 5건 + AAOS/TARA 매핑" — keep
    5: ("초기 계획서의 5 사례 카테고리(하드코딩 자격증명·권한 우회 4중 차단·평문 통신·exported ContentProvider·"
        "차량 제어 spoofing) 모두에 실측 finding 매핑. 가짜 수치 없이 자동 탐지/검증한 라벨에서만 발췌"),
    6: ("AAOS 보안 가이드라인 §3.7 권한 모델·§4.2 자격증명 보호·§5.1 통신 보안 자동 커버. ISO/SAE 21434 기반 "
        "TARA Risk Matrix 자동 생성 — Critical 2건·High 6건·Medium 4건·Low 2건"),
    7: ("Critical 2건 모두 차량 제어 자산 직격: GleoActionSender 의 차량 명령 implicit broadcast(spoofing 가능), "
        "VehicleBroadcastReceiver 의 exported macAddress 미검증(DB 주입 가능). 자산 = Severe + 공격 실현가능성 = High"),
    # 8: "2) " + "본 학기 연구 RQ 첫 corpus-level 측정 3종" — keep
    9: ("RQ1 단계별 verdict transition 추적: 19개 finding의 Stage 1→Stage 2(GT)→Stage 3(consensus) 변화. Stage 2 "
        "GT 오탐 4건이 Stage 3 합의에서 모두 정확 필터링(0건 잘못 promote). 단 n=19 표본 우연성 배제는 다음 단계"),
    10: ("RQ2 다중 시각 합의 진단: 시각 간 Cohen's κ 측정. 방어자 시각이 19/19 finding 모두 flag(universal) 발견. "
         "다양성은 공격자 ↔ 도메인 전문가 사이에서만 의미(κ=0.759). 방어자 prompt selectivity 보강"),
    11: ("RQ3 Java-only 정적 분석 boundary 정량: PleOS Connect 207 시스템 APK 중 10.6%가 native lib 보유. "
         "보안 가치 큰 컴포넌트(차량 제어·맵·음성 비서)가 native top — 한계 L1 의 위험도 가중 인덱스"),
    # 12: "3) " + "보고서 v0.9-r1 + 외부 reader 가독성 정제 + bootstrap CI" — keep
    13: ("보고서 표지 학번/소속 삭제, Notation & Glossary 신설(Stage 0/1/2/3 명명·Phase A~E·한계 L1~L5·Finding ID "
         "컨벤션). docs 8개 → 4개로 정리(reading order 번호 prefix). 영문/한국어 prompt 통일"),
    14: ("per-APK 보고서 evidence 마스킹 후 공개판 별도 디렉토리(data/reports/public/) — contractor IP 보호 정책 "
         "준수하면서 분석 메타데이터(class·line·category·rationale·매핑)는 그대로 유지로 외부 검증 가능"),
    15: ("향후 계획(11주차): src/eval.py 에 bootstrap CI(1000 iter) 추가 완료(95% CI Precision [57.9%, 94.7%] / "
         "F1 [73.3%, 97.3%], n=19 caveat 정량화). 표본 확장(n ≥ 30) + 한계+비용 1슬라이드"),
}


def update_paragraph(para, new_text: str) -> None:
    """Replace paragraph text in-place — preserves font / level / alignment of run 0."""
    if not para.runs:
        run = para.add_run()
        run.text = new_text
        return
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ""


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"PPT base not found: {SRC}")
    DST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SRC, DST)

    prs = Presentation(str(DST))
    sl = prs.slides[0]

    target = None
    for shp in sl.shapes:
        if shp.has_text_frame and shp.name == "Rectangle 3":
            target = shp
            break
    if target is None:
        target = next(s for s in sl.shapes if s.has_text_frame)

    paragraphs = list(target.text_frame.paragraphs)
    n_para = len(paragraphs)
    print(f"target shape: {target.name}, paragraphs: {n_para}")
    if n_para != 16:
        print(f"warning: expected 16 paragraphs, got {n_para} — applying best-effort")

    for idx, new in NEW_TEXTS.items():
        if idx >= n_para:
            print(f"  skip idx={idx} (out of range)")
            continue
        para = paragraphs[idx]
        if isinstance(new, tuple):
            num, title = new
            if len(para.runs) >= 2:
                para.runs[0].text = num
                para.runs[1].text = title
                for r in para.runs[2:]:
                    r.text = ""
            else:
                update_paragraph(para, num + title)
        else:
            update_paragraph(para, new)

    prs.save(str(DST))
    print(f"Wrote: {DST}")
    print(f"Size: {DST.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
