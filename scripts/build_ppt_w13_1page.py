#!/usr/bin/env python
"""13주차 progress report — 1 page version.

9/10/11/12주차 v3 동일 textbox layout (Rectangle 3, 16 paragraphs).
사용자 base PPT pptx/자율주행연구프로젝트1_13주차_*.pptx 사용.

작성자/제출 파일명은 `scripts/_author.json` 에서 로드.
"""
import json
import shutil
from pathlib import Path

from pptx import Presentation


def _load_author() -> dict:
    here = Path(__file__).parent
    p = here / "_author.json"
    if not p.exists():
        p = here / "_author.example.json"
    with p.open(encoding="utf-8") as f:
        return json.load(f)


AUTHOR = _load_author()
ROOT = Path(__file__).parent.parent.parent
WEEK = "13"
PPT_BASENAME = f"자율주행연구프로젝트1_{WEEK}주차"
SRC = ROOT / "pptx" / f"{PPT_BASENAME}_{AUTHOR['student_id']}_{AUTHOR['author_name']}.pptx"
DST = ROOT / "pptx" / "v3" / f"{PPT_BASENAME}_{AUTHOR['student_id']}_{AUTHOR['author_name']}_v3.pptx"


# 16 paragraphs, indexed 0~15.
NEW_TEXTS: dict[int, str | tuple[str, str]] = {
    # 0: Progress Report (period) — keep
    # 1: 프로젝트 제목 — keep
    2: "Stage 3 ensemble 표본 확장 (n=28) + R4 real-world 난독화 baseline 측정 + 보고서 v0.9-r1 마무리:",
    3: ("R1.d 표본 확장(InsecureBankv2 9 finding) 까지 Stage 3 multi-perspective consensus "
        "확장 — Stage 3 ≥2/3 합의 F1 0.966 → 0.979. 동시에 R4 첫 측정으로 "
        "real-world ProGuard'd OSS(NewPipe) 도입해 hand-crafted MASTG 와 PleOS 사이의 "
        "obfuscation density baseline 정량화."),
    # 4: "1) " + "Stage 3 multi-perspective consensus 표본 확장 (n=19 → n=28)" — keep
    5: ("InsecureBankv2 9 finding 에 attacker / defender / domain_expert 시각 3종 평가 적용. "
        "결과: 7 strong-TP (3/3) + 2 TP (2/3, ib2-8/9 — domain_expert 가 banking-only context "
        "에서 abstain). 모두 의도된 vuln 이라 in-scope 100% TP"),
    6: ("Stage 3 ≥2/3 합의 측정값 갱신 (n=28): Precision 100%, Recall 95.8%, F1 0.979 "
        "(이전 n=19: 0.966). FN 1건 = vc-7 (LOW hardening, 1/3 합의로 강등). "
        "≥3/3 strong-TP 19건 vs ≥2/3 reported 23건 = 4건 차이는 ssl-4 / ucl1-3 / ib2-8 / ib2-9"),
    7: ("R2.a perspective disagreement 재측정 (n=28): attacker ↔ domain_expert κ=0.619 "
        "(이전 n=19: 0.759, 표본 확장으로 다양성 약간 증가). defender 는 28/28 universal flag — "
        "selectivity-calibrated prompt 의 효과 검증은 다음 corpus 측정에서"),
    # 8: "2) " + "R4 — real-world commercial 난독화 baseline 측정" — keep
    9: ("외부 raw URL 다운로드(NewPipe v0.27.6, OSS, ProGuard 활성) → jadx 디컴파일 "
        "(2,520 비-framework 클래스) → entropy 측정. HIGH 난독화(composite ≥ 0.7) "
        "**42 클래스 = 1.7%**. mean composite 0.120, mean jadx_ratio 0.080"),
    10: ("Real-world floor 의 정량적 위치: NewPipe(1.7%) 이 hand-crafted MASTG(40~50%) 와 "
         "PleOS(0.2~0.4%) 의 사이. 초기 계획서가 가정한 두 극단 사이에 commercial OSS baseline 끼워넣음"),
    11: ("R4 의 Stage 0 LLM rename exact accuracy 측정은 ProGuard mapping file 부재로 partial. "
         "본 학기는 entropy density baseline 만 정량 입증, exact accuracy 는 Future Work "
         "(mapping 확보 또는 debug build 다운로드 후)"),
    # 12: "3) " + "보고서 v0.9-r1 마무리 + 산출물 정리" — keep
    13: ("docs/01_report.md § 0 (요약 측정값 표) / § 3.2 (Stage 1/3 표) / § 3.4 (난독화 + NewPipe) / "
         "§ 3.6.1 (R1.b bootstrap CI) 모두 n=28 + 2026-05-08 새 측정값으로 갱신"),
    14: ("학기 누적 산출물: combined GT n=28 + 외부 corpus 7 sample(MASTG 6 + InsecureBankv2 1) + "
         "Stage 3 28 entries + bootstrap CI + perspective agreement + stage transition + native lib "
         "inventory + 6 차트 + 사례 5건 + AAOS/TARA 매핑"),
    15: ("향후 계획(14주차): 발표 PPT + 라이브 데모 시나리오(3~5분, APK 1개 추출 → 디컴파일 → Stage 1) + "
         "백업 데모 영상 + 예상 Q&A. 보고서 v0.9-r1 → v1.0 (지도교수 추가 피드백 반영 시)"),
}


def update_paragraph(para, new_text: str) -> None:
    if not para.runs:
        run = para.add_run()
        run.text = new_text
        return
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ""


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"PPT base not found: {SRC}")
    DST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SRC, DST)

    prs = Presentation(str(DST))
    sl = prs.slides[0]

    target = None
    for shp in sl.shapes:
        if shp.has_text_frame and shp.name == "Rectangle 3":
            target = shp
            break
    if target is None:
        target = next(s for s in sl.shapes if s.has_text_frame)

    paragraphs = list(target.text_frame.paragraphs)
    n_para = len(paragraphs)
    print(f"target shape: {target.name}, paragraphs: {n_para}")
    if n_para != 16:
        print(f"warning: expected 16 paragraphs, got {n_para} -- applying best-effort")

    for idx, new in NEW_TEXTS.items():
        if idx >= n_para:
            print(f"  skip idx={idx} (out of range)")
            continue
        para = paragraphs[idx]
        if isinstance(new, tuple):
            num, title = new
            if len(para.runs) >= 2:
                para.runs[0].text = num
                para.runs[1].text = title
                for r in para.runs[2:]:
                    r.text = ""
            else:
                update_paragraph(para, num + title)
        else:
            update_paragraph(para, new)

    prs.save(str(DST))
    print(f"Wrote: {DST}")
    print(f"Size: {DST.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
