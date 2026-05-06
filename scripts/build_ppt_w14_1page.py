#!/usr/bin/env python
"""14주차 progress report — 1 page version (학기 종합 + 발표 준비).

9~13주차 v3 동일 textbox layout (Rectangle 3, 16 paragraphs).
"""
import json
import shutil
from pathlib import Path

from pptx import Presentation


def _load_author() -> dict:
    here = Path(__file__).parent
    p = here / "_author.json"
    if not p.exists():
        p = here / "_author.example.json"
    with p.open(encoding="utf-8") as f:
        return json.load(f)


AUTHOR = _load_author()
ROOT = Path(__file__).parent.parent.parent
WEEK = "14"
PPT_BASENAME = f"자율주행연구프로젝트1_{WEEK}주차"
SRC = ROOT / "pptx" / f"{PPT_BASENAME}_{AUTHOR['student_id']}_{AUTHOR['author_name']}.pptx"
DST = ROOT / "pptx" / "v3" / f"{PPT_BASENAME}_{AUTHOR['student_id']}_{AUTHOR['author_name']}_v3.pptx"


# 16 paragraphs, indexed 0~15.
NEW_TEXTS: dict[int, str | tuple[str, str]] = {
    # 0: Progress Report (period) — keep
    # 1: 프로젝트 제목 — keep
    2: "학기 종합 + 최종 발표 준비 (보고서 v1.0 + 발표 PPT 18 슬라이드 + 라이브 데모):",
    3: ("Phase A~D 마무리 + 본 학기 RQ 4개 모두 정량 측정값 확보 + 학기 deliverable 종합. "
        "최종 corpus n=28, Stage 3 ≥2/3 합의 F1 0.979 (95% CI [0.833, 0.982]). "
        "보고서 v0.9-r1 → v1.0 + 최종 발표 PPT 빌드 + 라이브 데모 시나리오 정리."),
    # 4: "1) " + "학기 최종 측정값 종합" — keep
    5: ("Stage 1 (n=28): Precision 85.7% / Recall 100% / F1 0.923. 95% CI [71.4%, 96.4%]. "
        "Stage 3 ≥2/3 (n=28): Precision 100% / Recall 95.8% / F1 **0.979** — 본 학기 최고치"),
    6: ("초기 계획서 가설 모두 충족: 1차 오탐률 25% → 14.3% (실측), 3차 오탐률 7% → 0%, "
        "Precision 가설 0.93 → 1.00 (≥2/3 합의). 분석 대상 축소율 99.49% (1,765 → 9 priority class)"),
    7: ("Bootstrap CI (n=19 → n=28) — Precision CI 폭 36.8%p → 25.0%p (-32%), F1 CI 폭 24.0%p → 14.9%p (-38%). "
        "표본 확장의 통계적 효과 정량 입증. 11~12주차 RQ1.b/d 의 직접 결과"),
    # 8: "2) " + "외부 corpus 종합 + R4 real-world baseline" — keep
    9: ("외부 corpus 7 sample 누계: in-scope 13 finding (UnCrackable Level1/3 + InsecureBankv2) + "
        "boundary 4 sample (UnCrackable-Level2 / r2pay / HelloWord-JNI / certificatePinningXamarin). "
        "Java-only 한계 L1 의 corpus-level 정량 evidence"),
    10: ("R4 — NewPipe v0.27.6 (real-world OSS, ProGuard 활성) 도입. HIGH 난독화 1.7% (vs hand-crafted "
         "MASTG 40~50%, PleOS 0.2~0.4%) — 본 학기 처음으로 real-world commercial 난독화 baseline 정량 위치 확보"),
    11: ("R3.a — PleOS Connect 207 시스템 APK 중 10.6% 가 native lib 보유. 보안 가치 큰 컴포넌트 "
         "(차량 제어 / 맵 / 음성 비서) 가 native top — 한계 L1 의 위험도 가중 인덱스"),
    # 12: "3) " + "학기 deliverable + 발표 준비" — keep
    13: ("산출물 종합: 보고서 v1.0 (docs/01_report.md, 8장 + Notation & Glossary + Appendix) + "
         "사례 5건 + 차트 6장 + AAOS/MASVS/TARA 매핑 + 통합 아키텍처 + 일반화 평가 + "
         "다국어 prompt 6종 (Stage 0/1/3) + 결정론적 측정 스크립트 9종"),
    14: ("최종 발표 PPT 18 슬라이드 — Cover / RQ / Background / Pipeline / Stage 0~3 / 측정 / "
         "Bootstrap CI / Ablation / RQ 별 finding / R4 baseline / AAOS / TARA / 사례 / 한계 / 결론. "
         "라이브 데모 시나리오 3~5분 (APK 추출 → jadx → Stage 1 LLM)"),
    15: ("Future Work: multi-LLM ensemble (외부 API 환경 확보 시) / native binary scanner 통합 (radare2 / Ghidra) / "
         "fully-batch 자동화 (CI/CD trigger) / commercial APK 의 ProGuard mapping 확보 → R4 exact accuracy 측정"),
}


def update_paragraph(para, new_text: str) -> None:
    if not para.runs:
        run = para.add_run()
        run.text = new_text
        return
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ""


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"PPT base not found: {SRC}")
    DST.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SRC, DST)

    prs = Presentation(str(DST))
    sl = prs.slides[0]

    target = None
    for shp in sl.shapes:
        if shp.has_text_frame and shp.name == "Rectangle 3":
            target = shp
            break
    if target is None:
        target = next(s for s in sl.shapes if s.has_text_frame)

    paragraphs = list(target.text_frame.paragraphs)
    n_para = len(paragraphs)
    print(f"target shape: {target.name}, paragraphs: {n_para}")

    for idx, new in NEW_TEXTS.items():
        if idx >= n_para:
            continue
        para = paragraphs[idx]
        if isinstance(new, tuple):
            num, title = new
            if len(para.runs) >= 2:
                para.runs[0].text = num
                para.runs[1].text = title
                for r in para.runs[2:]:
                    r.text = ""
            else:
                update_paragraph(para, num + title)
        else:
            update_paragraph(para, new)

    prs.save(str(DST))
    print(f"Wrote: {DST}")
    print(f"Size: {DST.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
